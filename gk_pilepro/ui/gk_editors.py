# -*- coding: utf-8 -*-

import copy
import tkinter as tk
from tkinter import messagebox, simpledialog, ttk

from openpyxl.utils import get_column_letter
from PIL import Image, ImageTk

from gk_pilepro.ui.gk_ui import (
    RoundedMappingDropdown,
    RoundedMappingLabel,
    UI_BORDER,
    UI_MUTED,
    UI_SURFACE,
    UI_TEXT,
    ui_button,
)

from gk_pilepro.gk_excel import (
    _static_jacking_to_float,
    norm,
    short_header_name,
)

class MappingEditor(tk.Frame):

    def __init__(self, parent):

        super().__init__(parent, bg=UI_SURFACE)

        self.pack(fill="x", pady=(8, 0))



        self.source_bg = "#ffffff"

        self.target_bg = "#f7fbf9"

        self.mapping_border = UI_BORDER

        self.combo_style = "Mapping.TCombobox"

        try:

            style = ttk.Style()

            style.configure(

                self.combo_style,

                fieldbackground=self.target_bg,

                background=self.target_bg,

                foreground=UI_TEXT,

                bordercolor=self.mapping_border,

                lightcolor=self.mapping_border,

                darkcolor=self.mapping_border,

                arrowcolor=UI_MUTED,

                padding=2,

            )

            style.map(

                self.combo_style,

                fieldbackground=[("readonly", self.target_bg)],

                background=[("readonly", self.target_bg)],

                selectbackground=[("readonly", "#cfe5ff")],

                selectforeground=[("readonly", UI_TEXT)],

            )

        except Exception:

            pass



        self.mapping_vars = []

        self.excel_headers = []

        self.table_cols = []

        self.auto_map_idx = []

        self.display_to_col = {}

        self.col_to_display = {}
        self.on_mapping_change = None

        self._layout_width = 0

        self._relayout_job = None



        self.title = tk.Label(

            self,

            text="Xác nhận mapping cột",

            font=("Segoe UI", 11, "bold"),

            bg=UI_SURFACE,

            fg=UI_TEXT,

        )

        self.title.pack(anchor="w")



        # Khung cuộn mapping

        self.box = tk.Frame(self, bg=UI_SURFACE)

        self.box.pack(fill="x", expand=False, pady=(6, 0))



        canvas_h = int(getattr(parent, "mapping_canvas_h", 260))

        self.canvas = tk.Canvas(

            self.box,

            height=canvas_h,

            highlightthickness=1,

            highlightbackground=UI_BORDER,

            bg=UI_SURFACE,

            bd=0,

        )

        self.scrollbar = ttk.Scrollbar(self.box, orient="vertical", command=self.canvas.yview)



        self.inner = tk.Frame(self.canvas, bg=UI_SURFACE)

        self.inner_id = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")



        self.canvas.configure(yscrollcommand=self.scrollbar.set)



        self.canvas.pack(side="left", fill="x", expand=True)

        self.scrollbar.pack(side="right", fill="y")



        self.inner.bind("<Configure>", self._on_frame_configure)

        self.canvas.bind("<Configure>", self._on_canvas_configure)



        self.canvas.bind("<Enter>", self._bind_mousewheel)

        self.canvas.bind("<Leave>", self._unbind_mousewheel)



    def _on_frame_configure(self, event=None):

        self.canvas.configure(scrollregion=self.canvas.bbox("all"))



    def _on_canvas_configure(self, event=None):

        try:

            self.canvas.itemconfig(self.inner_id, width=event.width)

        except Exception:

            pass

        try:
            width = int(event.width or 0)
        except Exception:
            width = 0

        if width <= 0:
            return

        if self.table_cols and self.excel_headers and abs(width - self._layout_width) >= 24:
            self._layout_width = width
            if self._relayout_job is not None:
                try:
                    self.after_cancel(self._relayout_job)
                except Exception:
                    pass
            selected_values = [var.get() for var in self.mapping_vars]
            self._relayout_job = self.after_idle(lambda: self._rebuild_for_width(width, selected_values))



    def _on_mousewheel(self, event):

        try:

            self.canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        except Exception:

            pass



    def _bind_mousewheel(self, event=None):

        try:

            self.canvas.bind_all("<MouseWheel>", self._on_mousewheel)

        except Exception:

            pass



    def _unbind_mousewheel(self, event=None):

        try:

            self.canvas.unbind_all("<MouseWheel>")

        except Exception:

            pass



    def _layout_width_hint(self):

        candidates = []

        for widget in (self.box, self.canvas, self.winfo_toplevel()):

            try:

                width = int(widget.winfo_width() or 0)

                if width > 0:

                    candidates.append(width)

            except Exception:

                pass

        if candidates:

            return max(candidates)

        try:

            return int(self.winfo_toplevel().winfo_screenwidth() or 1366)

        except Exception:

            return 1366



    def _compute_layout_sizes(self, available_width):

        usable = max(260, available_width - 18)

        stacked = usable < 420

        src_width = int(usable * (0.30 if stacked else 0.18))

        src_width = max(64, min(132 if stacked else 128, src_width))

        target_width = usable - src_width - 32

        if stacked:

            target_width = max(180, usable - 18)

        elif target_width < 140:

            src_width = max(80, min(src_width, usable - 172))

            target_width = usable - src_width - 32

        target_width = max(140 if not stacked else 180, min(360 if stacked else 320, target_width))

        src_max_len = max(5, min(14, int((src_width - 18) / 6)))

        row_pady = 2 if usable < 500 else 3

        return src_width, target_width, src_max_len, row_pady, stacked



    def _rebuild_for_width(self, width, selected_values=None):

        try:

            self.set_mapping(self.table_cols, self.excel_headers, self.auto_map_idx, selected_values=selected_values)

        finally:

            self._relayout_job = None



    def clear(self, notify=True):

        for w in self.inner.winfo_children():

            w.destroy()



        self.mapping_vars = []

        self.excel_headers = []

        self.table_cols = []

        self.display_to_col = {}

        self.col_to_display = {}



        try:

            self.canvas.yview_moveto(0)

        except Exception:

            pass
        if notify:
            self._notify_mapping_change()



    def _make_excel_choices(self, excel_headers):

        choices = ["(bỏ qua)"]

        used = set()

        self.display_to_col = {}

        self.col_to_display = {}



        for col_idx, name in excel_headers:

            try:

                short = short_header_name(name, 34)

            except Exception:

                short = str(name)



            label = f"{get_column_letter(col_idx)}: {short}"



            base = label

            n = 2

            while label in used:

                label = f"{base} ({n})"

                n += 1



            used.add(label)

            choices.append(label)

            self.display_to_col[label] = col_idx

            self.col_to_display[col_idx] = label



        return choices



    def set_mapping(self, table_cols, excel_headers, auto_map_idx, selected_values=None):

        if selected_values is None and getattr(self, "mapping_vars", None):
            try:
                selected_values = [var.get() for var in self.mapping_vars]
            except Exception:
                selected_values = []

        self.clear(notify=False)

        self.table_cols = list(table_cols or [])

        self.excel_headers = list(excel_headers or [])

        self.auto_map_idx = list(auto_map_idx or [])

        if selected_values is None:

            selected_values = []

        available_width = self._layout_width_hint()
        self._layout_width = available_width
        src_width, target_width, src_max_len, row_pady, stacked = self._compute_layout_sizes(available_width)



        if not self.table_cols or not self.excel_headers:
            msg = "Chưa có cột OCR hoặc cột Excel để mapping."
            if not self.table_cols and self.excel_headers:
                msg = "Chưa có cột OCR để mapping. Hãy đọc bảng/đọc phiếu cọc trước."
            elif self.table_cols and not self.excel_headers:
                msg = "Chưa có cột Excel để mapping. Hãy chọn Excel và đọc sheet trước."
            tk.Label(
                self.inner,
                text=msg,
                bg=UI_SURFACE,
                fg=UI_MUTED,
                font=("Segoe UI", 10),
                wraplength=max(220, self._layout_width_hint() - 28),
                justify="left",
            ).pack(fill="x", padx=8, pady=10)
            self.canvas.update_idletasks()
            self.canvas.configure(scrollregion=self.canvas.bbox("all"))
            self.canvas.yview_moveto(0)
            self._notify_mapping_change()
            return

        excel_choices = self._make_excel_choices(self.excel_headers)



        for i, src in enumerate(self.table_cols):

            row = tk.Frame(self.inner, bg=UI_SURFACE)

            row.pack(fill="x", pady=row_pady)



            try:

                src_show = short_header_name(src, src_max_len)

            except Exception:

                src_show = str(src)



            var = tk.StringVar()



            if i < len(selected_values) and selected_values[i] in excel_choices:

                var.set(selected_values[i])

            elif i >= len(auto_map_idx) or auto_map_idx[i] is None:

                var.set("(bỏ qua)")

            else:

                try:

                    excel_col_idx = self.excel_headers[auto_map_idx[i]][0]

                    var.set(self.col_to_display.get(excel_col_idx, "(bỏ qua)"))

                except Exception:

                    var.set("(bỏ qua)")



            if stacked:
                row.grid_columnconfigure(0, weight=1)
                src_box = RoundedMappingLabel(
                    row,
                    text=src_show,
                    bg_color=self.source_bg,
                    border_color=UI_BORDER,
                    width=src_width,
                    height=28,
                )
                src_box.grid(row=0, column=0, sticky="ew", padx=(1, 6), pady=(0, 4))

                mapping_line = tk.Frame(row, bg=UI_SURFACE)
                mapping_line.grid(row=1, column=0, sticky="ew", padx=(1, 6))
                mapping_line.grid_columnconfigure(0, weight=0)
                mapping_line.grid_columnconfigure(1, weight=1)

                tk.Label(
                    mapping_line,
                    text="→",
                    width=1,
                    anchor="center",
                    bg=UI_SURFACE,
                    fg=UI_MUTED,
                ).grid(row=0, column=0, sticky="w", padx=(0, 6))

                cb = RoundedMappingDropdown(
                    mapping_line,
                    values=excel_choices,
                    variable=var,
                    bg_color=self.target_bg,
                    border_color=self.mapping_border,
                    width=target_width,
                    height=30,
                )
                cb.grid(row=0, column=1, sticky="ew")
            else:
                row.grid_columnconfigure(0, weight=0, minsize=src_width)
                row.grid_columnconfigure(1, weight=0, minsize=14)
                row.grid_columnconfigure(2, weight=1)

                src_box = RoundedMappingLabel(
                    row,
                    text=src_show,
                    bg_color=self.source_bg,
                    border_color=UI_BORDER,
                    width=src_width,
                    height=28,
                )
                src_box.grid(row=0, column=0, sticky="w", padx=(1, 2))

                tk.Label(
                    row,
                    text="→",
                    width=1,
                    anchor="center",
                    bg=UI_SURFACE,
                    fg=UI_MUTED,
                ).grid(row=0, column=1, sticky="w", padx=(0, 1))

                cb = RoundedMappingDropdown(
                    row,
                    values=excel_choices,
                    variable=var,
                    bg_color=self.target_bg,
                    border_color=self.mapping_border,
                    width=target_width,
                    height=30,
                )
                cb.grid(row=0, column=2, sticky="ew", padx=(0, 6))



            self.mapping_vars.append(var)
            var.trace_add("write", self._on_mapping_var_changed)



        self.canvas.update_idletasks()

        self.canvas.configure(scrollregion=self.canvas.bbox("all"))

        self.canvas.yview_moveto(0)
        self._notify_mapping_change()



    def get_mapping(self):

        out = []

        for var in self.mapping_vars:

            chosen = var.get()

            if chosen == "(bỏ qua)":

                out.append(None)

            else:

                out.append(self.display_to_col.get(chosen))

        return out

    def get_mapping_stats(self):
        mapping = self.get_mapping()
        total = len(mapping)
        mapped = sum(1 for excel_col in mapping if excel_col is not None)
        return {
            "mapped": mapped,
            "unmapped": max(0, total - mapped),
            "total": total,
        }

    def _on_mapping_var_changed(self, *_args):
        self._notify_mapping_change()

    def _notify_mapping_change(self):
        callback = getattr(self, "on_mapping_change", None)
        if callable(callback):
            try:
                callback(self.get_mapping_stats())
            except Exception:
                pass


class TableEditor(tk.Frame):

    def __init__(self, parent):

        super().__init__(parent, bg=UI_SURFACE)

        self.pack(fill="both", expand=True)

        self.tables = []

        self.current = 0

        self.active_cell = None
        self.on_change = None
        self.on_selection_change = None
        self.on_compare_image = None
        self._selection_after_id = None
        self._undo_stack = []
        self._redo_stack = []
        self._history_limit = 50
        self._active_edit = None
        self.invalid_rows = set()



        top = tk.Frame(self, bg=UI_SURFACE)

        top.pack(fill="x", pady=(0, 8))

        top.grid_columnconfigure(1, weight=1)

        tk.Label(top, text="Bảng:", bg=UI_SURFACE, fg=UI_TEXT).grid(row=0, column=0, sticky="w", padx=(0, 8), pady=2)

        try:
            screen_w = int(parent.winfo_toplevel().winfo_screenwidth() or 1366)
        except Exception:
            screen_w = 1366
        combo_width = 320 if screen_w >= 1600 else 290 if screen_w >= 1400 else 260

        self.combo = RoundedMappingDropdown(

            top,

            values=[],

            variable=tk.StringVar(),

            bg_color="#f8fbff",

            border_color=UI_BORDER,

            width=combo_width,

            height=34,

            radius=8,

        )

        self.combo.grid(row=0, column=1, sticky="ew", padx=(0, 12), pady=2)

        self.combo.bind("<<ComboboxSelected>>", lambda e: self.switch_table())



        buttons = tk.Frame(top, bg=UI_SURFACE)

        buttons.grid(row=0, column=2, sticky="e", pady=0)

        ui_button(buttons, "Hoàn tác", self.undo, width=9, variant="soft").pack(side="left", padx=(0, 6))

        ui_button(buttons, "Làm lại", self.redo, width=8, variant="soft").pack(side="left", padx=(0, 8))

        ui_button(buttons, "Thêm dòng", self.add_row, width=11, variant="soft").pack(side="left", padx=(0, 8))

        ui_button(buttons, "Xóa dòng", self.delete_row, width=10).pack(side="left", padx=(0, 8))

        ui_button(buttons, "Sửa ô", self.edit_selected_cell, width=10).pack(side="left", padx=(0, 8))

        ui_button(buttons, "Xóa ô", self.clear_selected_cell, width=9).pack(side="left", padx=(0, 8))

        ui_button(
            buttons,
            "Đối chiếu ảnh",
            self.compare_selected_image,
            width=12,
            variant="soft",
        ).pack(side="left")

        self.source_crop_frame = tk.Frame(
            self,
            bg="#f8fbff",
            highlightthickness=1,
            highlightbackground=UI_BORDER,
        )
        self.source_crop_title = tk.Label(
            self.source_crop_frame,
            text="ẢNH DÒNG GỐC",
            bg="#f8fbff",
            fg=UI_MUTED,
            font=("Segoe UI", 9, "bold"),
            anchor="w",
        )
        self.source_crop_title.pack(fill="x", padx=8, pady=(5, 2))
        self.source_crop_label = tk.Label(
            self.source_crop_frame,
            bg="#ffffff",
            fg=UI_MUTED,
            text="Click một ô để hiện đúng dòng ảnh gốc tại đây.",
            font=("Segoe UI", 10),
            anchor="center",
            height=4,
        )
        self.source_crop_label.pack(fill="x", padx=8, pady=(0, 7))
        self.source_crop_frame.pack(fill="x", pady=(0, 8))
        self._source_crop_photo = None



        tree_frame = tk.Frame(self, bg=UI_SURFACE, highlightthickness=1, highlightbackground=UI_BORDER)

        tree_frame.pack(fill="both", expand=True)



        self.tree = ttk.Treeview(tree_frame, show="headings", style="Preview.Treeview")

        self.tree.tag_configure("preview_odd", background="#ffffff", foreground="#1f2933")

        self.tree.tag_configure("preview_even", background="#f7faf9", foreground="#1f2933")
        self.tree.tag_configure(
            "preview_invalid",
            background="#fff0f0",
            foreground="#b42318",
        )

        self.v_scroll = ttk.Scrollbar(tree_frame, orient="vertical", command=self._tree_yview, style="Vertical.TScrollbar")

        self.h_scroll = ttk.Scrollbar(tree_frame, orient="horizontal", command=self._tree_xview, style="Horizontal.TScrollbar")

        self.preview_grid_lines = []



        self.tree.configure(yscrollcommand=self._tree_yscroll, xscrollcommand=self.h_scroll.set)



        self.tree.grid(row=0, column=0, sticky="nsew")

        self.v_scroll.grid(row=0, column=1, sticky="ns")

        self.h_scroll.grid(row=1, column=0, sticky="ew")



        tree_frame.rowconfigure(0, weight=1)

        tree_frame.columnconfigure(0, weight=1)

        self.empty_preview = tk.Frame(tree_frame, bg="#ffffff")
        tk.Label(
            self.empty_preview,
            text="▦",
            bg="#ffffff",
            fg="#8aa39a",
            font=("Segoe UI", 34, "bold"),
        ).pack(anchor="center")
        tk.Label(
            self.empty_preview,
            text="Chưa có dữ liệu",
            bg="#ffffff",
            fg=UI_TEXT,
            font=("Segoe UI", 12, "bold"),
        ).pack(anchor="center", pady=(8, 0))
        tk.Label(
            self.empty_preview,
            text="Đọc dữ liệu từ Excel để hiển thị bảng tại đây",
            bg="#ffffff",
            fg=UI_MUTED,
            font=("Segoe UI", 10),
        ).pack(anchor="center", pady=(6, 0))
        self.empty_preview.place(relx=0.5, rely=0.5, anchor="center")



        self.tree.bind("<Double-1>", self.edit_cell)

        self.tree.bind("<ButtonRelease-1>", self.remember_active_cell)

        self.tree.bind("<Configure>", lambda _e: self._refresh_preview_grid())

        self.tree.bind("<F2>", lambda e: self.edit_selected_cell())

        self.tree.bind("<Return>", lambda e: self.edit_selected_cell())

        self.tree.bind("<Delete>", lambda e: self.clear_selected_cell())

        self.tree.bind("<MouseWheel>", self._on_tree_mousewheel)

        self.tree.bind("<Shift-MouseWheel>", self._on_tree_shift_mousewheel)
        self.tree.bind("<Control-z>", lambda _e: self.undo())
        self.tree.bind("<Control-Z>", lambda _e: self.undo())
        self.tree.bind("<Control-y>", lambda _e: self.redo())
        self.tree.bind("<Control-Y>", lambda _e: self.redo())



    def _on_tree_mousewheel(self, event):

        try:

            self.tree.yview_scroll(int(-1 * (event.delta / 120)), "units")

            self._refresh_preview_grid()

            return "break"

        except Exception:

            return None



    def _on_tree_shift_mousewheel(self, event):

        try:

            self.tree.xview_scroll(int(-1 * (event.delta / 120)), "units")

            self._refresh_preview_grid()

            return "break"

        except Exception:

            return None



    def _tree_xview(self, *args):

        self.tree.xview(*args)

        self._refresh_preview_grid()



    def _tree_yview(self, *args):

        self.tree.yview(*args)

        self._refresh_preview_grid()



    def _tree_yscroll(self, first, last):

        self.v_scroll.set(first, last)

        self._refresh_preview_grid()



    def _clear_preview_grid(self):

        for line in getattr(self, "preview_grid_lines", []):

            try:

                line.destroy()

            except Exception:

                pass

        self.preview_grid_lines = []



    def _refresh_preview_grid(self):

        try:

            cols = list(self.tree["columns"])

        except Exception:

            return

        self._clear_preview_grid()

        if not cols:

            return

        try:

            total_width = sum(int(self.tree.column(c, "width")) for c in cols)

            if total_width <= 0:

                return

            x_offset = int(float(self.tree.xview()[0]) * total_width)

            visible_width = self.tree.winfo_width()

            height = self.tree.winfo_height()

            x = 0

            for c in cols[:-1]:

                x += int(self.tree.column(c, "width"))

                line_x = x - x_offset

                if -2 <= line_x <= visible_width + 2:

                    line = tk.Frame(self.tree, bg="#e1e8e5", width=1, bd=0, highlightthickness=0)

                    line.place(x=line_x, y=0, width=1, height=height)

                    line.lift()

                    self.preview_grid_lines.append(line)



            for item in self.tree.get_children():

                bbox = self.tree.bbox(item)

                if not bbox:

                    continue

                _x, y, _w, row_h = bbox

                line_y = y + row_h - 1

                if 0 <= line_y <= height:

                    line = tk.Frame(self.tree, bg="#e8efec", height=1, bd=0, highlightthickness=0)

                    line.place(x=0, y=line_y, width=visible_width, height=1)

                    line.lift()

                    self.preview_grid_lines.append(line)

        except Exception:

            self._clear_preview_grid()



    def set_tables(self, tables):

        self.tables = tables or []
        self._undo_stack = []
        self._redo_stack = []

        self.current = 0

        names = [t.get("title") or f"Bảng {i+1}" for i, t in enumerate(self.tables)]

        self.combo["values"] = names

        if names:

            self.combo.current(0)

        self.render()
        self._notify_change()



    def switch_table(self):

        self.sync_current_from_tree()

        self.current = self.combo.current()

        self.render()
        self._notify_change()


    def _notify_change(self):

        cb = getattr(self, "on_change", None)

        if callable(cb):

            try:

                cb(self.get_tables())

            except TypeError:

                try:

                    cb()

                except Exception:

                    pass

            except Exception:

                pass



    def render(self):

        for item in self.tree.get_children():

            self.tree.delete(item)

        self.tree["columns"] = []



        if not self.tables:
            try:
                self.empty_preview.place(relx=0.5, rely=0.5, anchor="center")
            except Exception:
                pass

            return

        try:
            self.empty_preview.place_forget()
        except Exception:
            pass



        t = self.tables[self.current]

        cols = t["columns"]

        self.tree["columns"] = cols



        for c in cols:

            self.tree.heading(c, text=c)

            self.tree.column(c, width=max(100, min(190, len(str(c)) * 10)), anchor="center", stretch=False)



        for idx, row in enumerate(t["rows"]):

            rr = row[:len(cols)] + [""] * max(0, len(cols) - len(row))

            tag = (
                "preview_invalid"
                if idx + 1 in self.invalid_rows
                else ("preview_even" if idx % 2 else "preview_odd")
            )

            self.tree.insert("", "end", values=rr, tags=(tag,))

        self.tree.after_idle(self._refresh_preview_grid)

    def set_invalid_rows(self, row_numbers):
        self.invalid_rows = {
            int(row)
            for row in (row_numbers or [])
            if str(row).strip().isdigit() and int(row) > 0
        }
        for idx, item in enumerate(self.tree.get_children()):
            tag = (
                "preview_invalid"
                if idx + 1 in self.invalid_rows
                else ("preview_even" if idx % 2 else "preview_odd")
            )
            self.tree.item(item, tags=(tag,))



    def sync_current_from_tree(self):

        if not self.tables:

            return

        active_edit = getattr(self, "_active_edit", None)
        if isinstance(active_edit, dict):
            entry = active_edit.get("entry")
            item = active_edit.get("item")
            idx = active_edit.get("index")
            try:
                if (
                    entry is not None
                    and entry.winfo_exists()
                    and item in self.tree.get_children()
                    and isinstance(idx, int)
                    and idx >= 0
                ):
                    values = list(self.tree.item(item, "values"))
                    values += [""] * (idx + 1 - len(values))
                    values[idx] = entry.get()
                    self.tree.item(item, values=values)
                    self._recalculate_combination_length(item, idx)
            except Exception:
                pass

        rows = []

        for item in self.tree.get_children():

            rows.append(list(self.tree.item(item, "values")))

        self.tables[self.current]["rows"] = rows


    def _recalculate_combination_length(self, item, changed_index):

        columns = list(self.tree["columns"] or [])
        if not columns or changed_index < 0 or changed_index >= len(columns):

            return False

        segment_names = {f"d{number}" for number in range(1, 7)}
        if norm(columns[changed_index]) not in segment_names:

            return False

        target_names = {
            "chieu dai to hop coc",
            "tong chieu dai to hop coc",
            "chieu dai coc",
            "pile combination length",
        }
        target_index = next(
            (
                index
                for index, name in enumerate(columns)
                if norm(name) in target_names
            ),
            None,
        )
        if target_index is None:

            return False

        values = list(self.tree.item(item, "values"))
        values += [""] * max(0, len(columns) - len(values))
        total = 0.0
        has_segment = False

        for index, name in enumerate(columns):
            if norm(name) not in segment_names:
                continue
            raw_value = str(values[index] or "").strip()
            if not raw_value:
                continue
            number = _static_jacking_to_float(raw_value)
            if number is None:
                return False
            total += number
            has_segment = True

        if not has_segment:
            total_value = ""
        elif abs(total - round(total)) < 1e-9:
            total_value = str(int(round(total)))
        else:
            total_value = (
                f"{total:.3f}".rstrip("0").rstrip(".").replace(".", ",")
            )

        if str(values[target_index] or "") == total_value:

            return False

        values[target_index] = total_value
        self.tree.item(item, values=values)
        return True



    def get_tables(self):

        self.sync_current_from_tree()

        return self.tables



    def get_current_table(self):

        self.sync_current_from_tree()

        if not self.tables:

            return None

        return self.tables[self.current]


    def _history_snapshot(self):

        self.sync_current_from_tree()

        return {
            "tables": copy.deepcopy(self.tables),
            "current": self.current,
        }


    def _push_undo(self):

        if not self.tables:

            return

        self._undo_stack.append(self._history_snapshot())
        if len(self._undo_stack) > self._history_limit:
            del self._undo_stack[0]
        self._redo_stack = []


    def _restore_history_snapshot(self, snapshot):

        self.tables = copy.deepcopy(snapshot.get("tables") or [])
        self.current = max(
            0,
            min(int(snapshot.get("current") or 0), max(0, len(self.tables) - 1)),
        )
        names = [
            table.get("title") or f"Bảng {idx + 1}"
            for idx, table in enumerate(self.tables)
        ]
        self.combo["values"] = names
        if names:
            self.combo.current(self.current)
        self.active_cell = None
        self.render()
        self._notify_change()


    def undo(self):

        if not self._undo_stack:

            return "break"

        self._redo_stack.append(self._history_snapshot())
        self._restore_history_snapshot(self._undo_stack.pop())
        return "break"


    def redo(self):

        if not self._redo_stack:

            return "break"

        self._undo_stack.append(self._history_snapshot())
        self._restore_history_snapshot(self._redo_stack.pop())
        return "break"



    def add_row(self):

        if not self.tables:

            return

        self._push_undo()
        cols = self.tables[self.current]["columns"]

        idx = len(self.tree.get_children())

        tag = "preview_even" if idx % 2 else "preview_odd"

        self.tree.insert("", "end", values=[""] * len(cols), tags=(tag,))
        row_sources = self.tables[self.current].setdefault("_row_source_indexes", [])
        row_sources.append(None)
        self.tables[self.current].setdefault("_row_bboxes", []).append(None)
        self.tables[self.current].setdefault("_cell_bboxes", []).append(
            [None] * len(cols)
        )

        self.tree.after_idle(self._refresh_preview_grid)

        self.sync_current_from_tree()
        self._notify_change()



    def delete_row(self):

        selected = list(self.tree.selection())
        if not selected:
            return
        self._push_undo()
        children = list(self.tree.get_children())
        deleted_indexes = sorted(
            (children.index(item) for item in selected if item in children),
            reverse=True,
        )

        for item in selected:

            self.tree.delete(item)

        row_sources = self.tables[self.current].get("_row_source_indexes")
        if isinstance(row_sources, list):
            for row_index in deleted_indexes:
                if 0 <= row_index < len(row_sources):
                    row_sources.pop(row_index)
        for metadata_key in ("_row_bboxes", "_cell_bboxes"):
            metadata = self.tables[self.current].get(metadata_key)
            if isinstance(metadata, list):
                for row_index in deleted_indexes:
                    if 0 <= row_index < len(metadata):
                        metadata.pop(row_index)

        self.sync_current_from_tree()
        self._notify_change()



    def remember_active_cell(self, event):

        item = self.tree.identify_row(event.y)

        col = self.tree.identify_column(event.x)

        if item and col:

            self.active_cell = (item, col)
            callback = getattr(self, "on_selection_change", None)
            if callable(callback):
                try:
                    if self._selection_after_id is not None:
                        self.after_cancel(self._selection_after_id)
                    context = self.get_selected_context()
                    self._selection_after_id = self.after(
                        240,
                        lambda: self._dispatch_selection_change(callback, context),
                    )
                except Exception:
                    pass


    def _dispatch_selection_change(self, callback, context):

        self._selection_after_id = None

        try:

            callback(context)

        except Exception:

            pass


    def get_selected_context(self):

        if not self.active_cell:

            return None

        item, col = self.active_cell
        children = list(self.tree.get_children())

        if item not in children:

            return None

        row_index = children.index(item)
        col_index = max(0, int(str(col).replace("#", "")) - 1)
        columns = list(self.tree["columns"] or [])
        values = list(self.tree.item(item, "values"))

        return {
            "table_index": self.current,
            "row_index": row_index,
            "column_index": col_index,
            "column": columns[col_index] if col_index < len(columns) else "",
            "value": values[col_index] if col_index < len(values) else "",
        }


    def compare_selected_image(self):

        callback = getattr(self, "on_compare_image", None)

        if callable(callback):

            try:

                callback(self.get_selected_context())

            except Exception:

                pass


    def show_source_row_crop(self, image_path, bbox, context=None):

        if not image_path or not bbox or len(bbox) != 4:
            self._source_crop_photo = None
            self.source_crop_label.configure(
                image="",
                text="Dòng này chưa có tọa độ ảnh. Hãy OCR lại ảnh bằng bản mới.",
                height=4,
            )
            return

        try:
            source = Image.open(image_path).convert("RGB")
            x1, y1, x2, y2 = [float(value) for value in bbox]
            left = int(source.width * x1 / 1000.0)
            top = int(source.height * y1 / 1000.0)
            right = int(source.width * x2 / 1000.0)
            bottom = int(source.height * y2 / 1000.0)

            row_height = max(1, bottom - top)
            pad_x = max(12, int(source.width * 0.012))
            pad_y = max(4, int(row_height * 0.25))
            crop_box = (
                max(0, left - pad_x),
                max(0, top - pad_y),
                min(source.width, right + pad_x),
                min(source.height, bottom + pad_y),
            )
            crop = source.crop(crop_box)

            self.update_idletasks()
            display_width = max(420, self.winfo_width() - 20)
            display_height = max(80, min(180, int(display_width * crop.height / max(1, crop.width))))
            crop.thumbnail((display_width, display_height), Image.Resampling.LANCZOS)
            self._source_crop_photo = ImageTk.PhotoImage(crop)

            row_number = int((context or {}).get("row_index", 0)) + 1
            column = str((context or {}).get("column") or "").strip()
            title = f"ẢNH DÒNG GỐC - Dòng {row_number}"
            if column:
                title += f" - Ô {column}"
            self.source_crop_title.configure(text=title)
            self.source_crop_label.configure(
                image=self._source_crop_photo,
                text="",
                height=0,
            )
        except Exception as exc:
            self._source_crop_photo = None
            self.source_crop_label.configure(
                image="",
                text=f"Không cắt được ảnh dòng gốc: {exc}",
                height=4,
            )



    def edit_selected_cell(self):

        if self.active_cell:

            item, col = self.active_cell

        else:

            selected = self.tree.selection()

            if not selected:

                return

            item = selected[0]

            col = "#1"

        self._edit_cell_by_item_col(item, col)



    def clear_selected_cell(self):

        if self.active_cell:

            item, col = self.active_cell

        else:

            selected = self.tree.selection()

            if not selected:

                return

            item = selected[0]

            col = "#1"



        if not item or not col:

            return



        idx = int(col.replace("#", "")) - 1

        vals = list(self.tree.item(item, "values"))

        vals += [""] * (idx + 1 - len(vals))

        if not vals[idx]:
            return
        self._push_undo()
        vals[idx] = ""

        self.tree.item(item, values=vals)
        self._recalculate_combination_length(item, idx)

        self.sync_current_from_tree()
        self._notify_change()



    def edit_cell(self, event):

        try:

            if self._selection_after_id is not None:

                self.after_cancel(self._selection_after_id)

                self._selection_after_id = None

        except Exception:

            pass

        item = self.tree.identify_row(event.y)

        col = self.tree.identify_column(event.x)

        if not item or not col:

            return

        self.active_cell = (item, col)

        self._edit_cell_by_item_col(item, col)



    def _edit_cell_by_item_col(self, item, col):

        idx = int(col.replace("#", "")) - 1

        bbox = self.tree.bbox(item, col)

        if not bbox:

            self.tree.see(item)

            bbox = self.tree.bbox(item, col)

            if not bbox:

                return



        x, y, w, h = bbox

        vals = list(self.tree.item(item, "values"))

        vals += [""] * (idx + 1 - len(vals))

        old = vals[idx]



        edit_var = tk.StringVar(value=str(old))
        ent = tk.Entry(self.tree, textvariable=edit_var)
        self._active_edit = {
            "entry": ent,
            "item": item,
            "index": idx,
        }

        ent.place(x=x, y=y, width=w, height=h)

        ent.focus()

        ent.select_range(0, "end")

        edit_state = {
            "done": False,
            "history_pushed": False,
            "last_value": str(old),
        }


        def apply_value(new_value):

            new_value = str(new_value)
            if new_value == edit_state["last_value"]:

                return

            if not edit_state["history_pushed"]:

                self._push_undo()
                edit_state["history_pushed"] = True

            vals2 = list(self.tree.item(item, "values"))

            vals2 += [""] * (idx + 1 - len(vals2))

            vals2[idx] = new_value

            self.tree.item(item, values=vals2)
            self._recalculate_combination_length(item, idx)

            edit_state["last_value"] = new_value
            self.sync_current_from_tree()
            self._notify_change()


        def save(_=None):

            if edit_state["done"]:

                return

            apply_value(ent.get())
            edit_state["done"] = True
            self._active_edit = None

            ent.destroy()



        def cancel(_=None):

            if edit_state["done"]:

                return

            edit_state["done"] = True
            self._active_edit = None
            if edit_state["last_value"] != str(old):

                vals2 = list(self.tree.item(item, "values"))
                vals2 += [""] * (idx + 1 - len(vals2))
                vals2[idx] = old
                self.tree.item(item, values=vals2)
                self._recalculate_combination_length(item, idx)
                self.sync_current_from_tree()
                self._notify_change()
                if edit_state["history_pushed"] and self._undo_stack:
                    self._undo_stack.pop()

            ent.destroy()



        ent.bind("<Return>", save)

        ent.bind("<FocusOut>", save)

        ent.bind("<Escape>", cancel)
        edit_var.trace_add("write", lambda *_args: apply_value(edit_var.get()))
